from __future__ import annotations

import asyncio
import re
import sys
from pathlib import Path

from openpyxl import load_workbook
from app.services.embeddings import embed_text, plant_embedding_text


def normalize(value: object) -> str:
    return re.sub(
        r"[^a-z0-9]+",
        " ",
        str(value or "").strip().lower()
    ).strip()


ALIASES = {
    "external_id": {
        "id",
        "plant id",
        "plant_id",
    },

    "sinhala_name": {
        "sinhala name",
        "sinhala local name",
        "local name",
        "sinhala_name",
    },

    "common_name": {
        "english name",
        "common name",
        "english_name",
        "common_name",
    },

    "scientific_name": {
        "scientific name",
        "scientific name as reported",
        "scientific_name",
    },

    "part_used": {
        "part",
        "part used",
        "parts used",
        "part_used",
    },

    "traditional_use": {
        "traditional use",
        "reported traditional use",
        "traditional_use",
        "reported use",
        "use",
    },

    "preparation": {
        "preparation",
        "traditional preparation",
        "traditional preparation reported",
        "reported preparation",
        "preparation method",
    },

    "safety_note": {
        "warning",
        "warnings",
        "safety",
        "safety note",
        "safety information",
    },

    "source_url": {
        "source url",
        "primary source url",
        "source",
        "reference",
        "url",
    },

    "verification_status": {
        "verification status",
        "verified",
        "verification",
    },

    "safety_review_status": {
        "safety review status",
        "safety status",
        "safety review",
    },
}


def find_column(
    headers: list[str],
    aliases: set[str]
) -> int | None:

    normalized_headers = [
        normalize(header)
        for header in headers
    ]

    # Exact match
    for alias in aliases:
        normalized_alias = normalize(alias)

        if normalized_alias in normalized_headers:
            return normalized_headers.index(
                normalized_alias
            )

    # Partial match
    for index, header in enumerate(
        normalized_headers
    ):
        for alias in aliases:
            normalized_alias = normalize(alias)

            if (
                normalized_alias
                and normalized_alias in header
            ):
                return index

    return None


def find_header_row(sheet) -> tuple[int, list[str]]:
    """
    Search the first 20 rows and find the real
    table header automatically.
    """

    for row_number, row in enumerate(
        sheet.iter_rows(
            min_row=1,
            max_row=20,
            values_only=True
        ),
        start=1,
    ):
        headers = [
            str(cell or "")
            for cell in row
        ]

        normalized_headers = [
            normalize(header)
            for header in headers
        ]

        has_id = "id" in normalized_headers

        has_traditional_use = any(
            "traditional use" in header
            for header in normalized_headers
        )

        if has_id and has_traditional_use:
            return row_number, headers

    raise RuntimeError(
        "Could not automatically find "
        "the plant dataset header row."
    )


def get_value(
    row: tuple,
    index: int | None
) -> str | None:

    if index is None:
        return None

    if index >= len(row):
        return None

    raw = row[index]

    if raw is None:
        return None

    text = str(raw).strip()

    return text or None


def is_verified(
    verification_status: str | None
) -> bool:

    status = normalize(
        verification_status
    )

    if not status or status.startswith("unverified") or "not verified" in status:
        return False

    return (
        status in {
            "yes",
            "true",
            "approved",
            "reviewed",
            "verified",
        }
        or "source verified" in status
        or status.startswith("verified ")
    )


async def run(path: Path) -> None:

    # Database-specific imports stay inside the import routine so the
    # workbook parsing helpers can be tested without a running database.
    from sqlalchemy import select

    from app.db.session import AsyncSessionLocal
    from app.models.plant import Plant

    print(
        f"Reading dataset: {path}"
    )

    workbook = load_workbook(
        path,
        read_only=True,
        data_only=True,
    )

    if "Plant Dataset" in workbook.sheetnames:
        sheet = workbook["Plant Dataset"]
    else:
        sheet = workbook.active

    # --------------------------------
    # Find real header row
    # --------------------------------

    header_row_number, headers = (
        find_header_row(sheet)
    )

    print(
        f"Header row detected: "
        f"{header_row_number}"
    )

    print(
        "Headers:",
        headers
    )

    # --------------------------------
    # Map columns
    # --------------------------------

    columns = {
        name: find_column(
            headers,
            aliases
        )
        for name, aliases
        in ALIASES.items()
    }

    print(
        "Detected columns:",
        columns
    )

    if columns["traditional_use"] is None:
        raise RuntimeError(
            "Could not find "
            "Traditional Use column."
        )

    # --------------------------------
    # Start reading AFTER header
    # --------------------------------

    rows = sheet.iter_rows(
        min_row=header_row_number + 1,
        values_only=True,
    )

    imported = 0
    updated = 0
    skipped = 0

    async with AsyncSessionLocal() as db:

        for row in rows:

            traditional_use = get_value(
                row,
                columns["traditional_use"],
            )

            # Ignore empty rows
            if not traditional_use:
                skipped += 1
                continue

            external_id = get_value(
                row,
                columns["external_id"],
            )

            scientific_name = get_value(
                row,
                columns["scientific_name"],
            )

            # --------------------------------
            # Check existing record
            # --------------------------------

            existing = None

            if external_id:
                result = await db.execute(
                    select(Plant).where(
                        Plant.external_id
                        == external_id
                    )
                )

                existing = (
                    result.scalar_one_or_none()
                )

            elif scientific_name:
                result = await db.execute(
                    select(Plant).where(
                        Plant.scientific_name
                        == scientific_name
                    )
                )

                existing = (
                    result.scalar_one_or_none()
                )

            verification_status = (
                get_value(
                    row,
                    columns[
                        "verification_status"
                    ],
                )
                or "UNVERIFIED"
            )

            payload = {

                "external_id":
                    external_id,

                "sinhala_name":
                    get_value(
                        row,
                        columns[
                            "sinhala_name"
                        ],
                    ),

                "common_name":
                    get_value(
                        row,
                        columns[
                            "common_name"
                        ],
                    ),

                "scientific_name":
                    scientific_name,

                "part_used":
                    get_value(
                        row,
                        columns[
                            "part_used"
                        ],
                    ),

                "traditional_use":
                    traditional_use,

                "preparation":
                    get_value(
                        row,
                        columns[
                            "preparation"
                        ],
                    ),

                "safety_note":
                    get_value(
                        row,
                        columns[
                            "safety_note"
                        ],
                    ),

                "source_url":
                    get_value(
                        row,
                        columns[
                            "source_url"
                        ],
                    ),

                "verification_status":
                    verification_status,

                "safety_review_status":
                    (
                        get_value(
                            row,
                            columns[
                                "safety_review_status"
                            ],
                        )
                        or "REVIEW_REQUIRED"
                    ),

                "is_verified":
                    is_verified(
                        verification_status
                    ),
            }

            # --------------------------------
            # Update or insert
            # --------------------------------

            if existing:

                for key, value in (
                    payload.items()
                ):
                    setattr(
                        existing,
                        key,
                        value,
                    )

                existing.embedding = embed_text(plant_embedding_text(existing))
                updated += 1

            else:

                plant = Plant(**payload)
                plant.embedding = embed_text(plant_embedding_text(plant))
                db.add(plant)

                imported += 1

        await db.commit()

    print()
    print("===========================")
    print("HelaCare Dataset Import")
    print("===========================")
    print(f"Imported : {imported}")
    print(f"Updated  : {updated}")
    print(f"Skipped  : {skipped}")
    print("===========================")


if __name__ == "__main__":

    if len(sys.argv) != 2:

        raise SystemExit(
            "Usage: "
            "python -m "
            "app.scripts.import_plants "
            "/path/to/workbook.xlsx"
        )

    asyncio.run(
        run(
            Path(
                sys.argv[1]
            )
        )
    )